from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
BASE_DIR = PROJECT_DIR.parent
OUTPUT_DIR = BASE_DIR / "output" / "doc"
ASSETS_DIR = OUTPUT_DIR / "assets"
FINAL_DOC_PATH = OUTPUT_DIR / "SocialSphere_lab6_report.docx"
FINAL_PDF_PATH = OUTPUT_DIR / "SocialSphere_lab6_report.pdf"
COMMIT_LINK_PLACEHOLDER = "https://github.com/getdhed/SocialSphere/compare/master...feature/async-data-storage"


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    return doc


def add_paragraph(
    doc: Document,
    text: str = "",
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_indent: bool = True,
    bold: bool = False,
    italic: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_indent:
        paragraph_format.first_line_indent = Cm(1.25)

    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)
    run.bold = bold
    run.italic = italic


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph_format.first_line_indent = Cm(0)

    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(11)


def add_heading(doc: Document, text: str) -> None:
    add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, bold=True)


def add_center(doc: Document, text: str, bold: bool = False) -> None:
    add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=bold)


def add_blank(doc: Document, count: int = 1) -> None:
    for _ in range(count):
        add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 15.6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    add_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)


def add_title_page(doc: Document) -> None:
    add_center(doc, "Министерство образования Республики Беларусь")
    add_blank(doc, 1)
    add_center(doc, "Учреждение образования")
    add_center(doc, "«БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ")
    add_center(doc, "ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ»")
    add_blank(doc, 1)
    add_paragraph(doc, "Факультет информационных технологий и управления", align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    add_paragraph(
        doc,
        "Кафедра информационных технологий автоматизированных систем",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_indent=False,
    )
    add_blank(doc, 7)
    add_center(doc, "По лабораторной работе №6")
    add_center(doc, "Вариант №5")
    add_center(doc, "«РАБОТА С WEB-API БРАУЗЕРА И АСИНХРОННЫМИ ОПЕРАЦИЯМИ»")
    add_blank(doc, 5)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    table.cell(0, 0).text = "Выполнили:"
    table.cell(0, 1).text = "ст. гр. 320601\nН.Д. Жебрун\nЮ.А. Гич"

    add_blank(doc, 1)
    table2 = doc.add_table(rows=1, cols=2)
    table2.autofit = True
    table2.cell(0, 0).text = "Проверил:"
    table2.cell(0, 1).text = "Н. В. Хаджинова"

    add_blank(doc, 6)
    add_center(doc, "Минск 2026")
    doc.add_page_break()


def fill_content(doc: Document) -> None:
    add_heading(doc, "1 ЦЕЛЬ РАБОТЫ")
    add_paragraph(
        doc,
        "Целью лабораторной работы является освоение асинхронного программирования в JavaScript, практическая работа с Fetch, "
        "обработкой JSON и организацией клиентского хранения данных в веб-приложении SocialSphere.",
    )
    add_paragraph(
        doc,
        "В рамках варианта №5 необходимо реализовать интеграцию с внешним API, автосохранение черновиков постов, "
        "очередь отправки данных при оффлайн-режиме и синхронизацию после восстановления соединения.",
    )

    add_heading(doc, "2 КРАТКИЕ ТЕОРЕТИЧЕСКИЕ СВЕДЕНИЯ")
    add_paragraph(
        doc,
        "Асинхронная модель JavaScript основана на цикле событий и неблокирующем выполнении операций ввода-вывода. "
        "Promise позволяет представить результат асинхронной операции в виде объекта, а конструкция async/await упрощает "
        "чтение и сопровождение кода, сохраняя последовательный стиль записи.",
    )
    add_paragraph(
        doc,
        "Fetch API используется для HTTP-запросов к REST-сервисам и возвращает Promise с объектом Response. "
        "На практике необходимо учитывать сетевые сбои, HTTP-статусы, таймауты, повторные попытки и отмену запросов через AbortController.",
    )
    add_paragraph(
        doc,
        "JSON является основным форматом обмена данными между клиентом и сервером. "
        "Для сохранения состояния на стороне клиента используются LocalStorage и SessionStorage: "
        "первое обеспечивает долговременное хранение, второе — хранение в рамках сессии вкладки.",
    )
    add_paragraph(
        doc,
        "Для оффлайн-функциональности применяется локальный кэш и очередь отправки изменений. "
        "После появления сети данные очереди отправляются на сервер, а интерфейс синхронизируется с актуальным состоянием.",
    )
    add_paragraph(
        doc,
        "С точки зрения безопасности ключи API не должны считаться защищенными в клиентском коде. "
        "Рекомендуется использовать HTTPS, ограничения ключей, OAuth и серверный прокси для работы с чувствительными данными.",
    )

    add_heading(doc, "3 ХОД РАБОТЫ")
    add_heading(doc, "3.1 Подготовка структуры проекта")
    add_paragraph(
        doc,
        "Проект SocialSphere был дополнен модульной структурой для работы с API, кэшем и парсингом данных. "
        "В директории js были созданы папки api и storage, добавлены отдельные сервисы localStorage/sessionStorage "
        "и модуль dataParser для нормализации данных, полученных от внешнего сервиса.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure1_structure.png",
        "Рисунок 1 – Обновленная структура проекта SocialSphere для лабораторной работы №6",
    )

    add_heading(doc, "3.2 Реализация API-сервиса и асинхронных запросов")
    add_paragraph(
        doc,
        "Для интеграции использован сервис JSONPlaceholder. В модуле apiService реализованы запросы GET и POST, "
        "обработка ошибок, таймаут через AbortController, ограничение методов и retry-логика для временных сбоев сети.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure2_api_service.png",
        "Рисунок 2 – Реализация API-сервиса (Fetch, retry, обработка ошибок)",
    )

    add_paragraph(doc, "Пример запроса и ответа API:")
    add_code(doc, "GET https://jsonplaceholder.typicode.com/posts?_start=0&_limit=15")
    add_code(doc, "Ответ 200 OK: [{ id: 1, userId: 1, title: \"...\", body: \"...\" }, ...]")
    add_code(doc, "POST https://jsonplaceholder.typicode.com/posts")
    add_code(doc, "Тело: { title: \"Новый пост\", body: \"Текст\", userId: 1 }")
    add_code(doc, "Ответ 201 Created: { id: 101, title: \"Новый пост\", body: \"Текст\", userId: 1 }")

    add_heading(doc, "3.3 Работа с клиентским хранилищем")
    add_paragraph(
        doc,
        "В локальном хранилище сохраняются кэш ленты, очередь отправки, тема и черновик поста. "
        "В сессионном хранилище сохраняются временные данные интерфейса: текущий черновик и список отображаемых постов. "
        "Данные записываются в формате JSON через отдельные сервисные модули.",
    )
    add_paragraph(
        doc,
        "Для поддержки консистентности добавлено версионирование хранилища. "
        "При изменении версии автоматически очищается устаревший кэш, что предотвращает отображение неактуальных данных.",
    )

    add_heading(doc, "3.4 Реализация индивидуального задания варианта №5")
    add_paragraph(
        doc,
        "Для выполнения индивидуального задания реализованы черновики постов с автосохранением и очередь отправки. "
        "Если пользователь находится оффлайн, пост отображается в ленте локально со статусом «В очереди», "
        "после события online запускается синхронизация и статус меняется на «Синхронизирован».",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure3_queue_sync.png",
        "Рисунок 3 – Логика очереди отправки и синхронизации постов",
    )
    add_paragraph(
        doc,
        "После демонстрации логики очереди важно отметить, что синхронизация запускается автоматически при появлении сети, "
        "а индикаторы состояния обновляются без перезагрузки страницы.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure7_interface.png",
        "Рисунок 4 – Интерфейс SocialSphere с индикаторами сети, очереди и кэша",
    )

    add_heading(doc, "3.5 Тестирование и обработка ошибок")
    add_paragraph(
        doc,
        "Проведены тесты в онлайн- и оффлайн-сценариях: загрузка ленты, публикация поста, автосохранение черновика, "
        "восстановление очереди после перезапуска и повторная синхронизация при восстановлении соединения.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure4_network.png",
        "Рисунок 5 – Сетевые запросы в DevTools (GET/POST, статусы 200/201, retry)",
    )
    add_paragraph(
        doc,
        "По сетевым логам видно, что приложение корректно обрабатывает статусы HTTP и выполняет повторные попытки при "
        "временных сбоях, после чего синхронизирует локальное состояние с сервером.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure5_storage_console.png",
        "Рисунок 6 – Данные LocalStorage/SessionStorage и логи асинхронных операций в консоли",
    )
    add_paragraph(
        doc,
        "Данные во вкладках Application и Console подтверждают корректное сохранение кэша, черновика и очереди отправки, "
        "что обеспечивает устойчивую работу интерфейса при нестабильном соединении.",
    )
    add_figure(
        doc,
        ASSETS_DIR / "lab6_figure6_offline.png",
        "Рисунок 7 – Поведение приложения при недоступности API (оффлайн и очередь отправки)",
    )

    add_paragraph(
        doc,
        "Ссылка на коммит с выполненной лабораторной работой: "
        f"{COMMIT_LINK_PLACEHOLDER} (после отправки ветки на GitHub).",
    )

    add_heading(doc, "4 ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ")
    qa_items = [
        (
            "1. Каковы основные различия между LocalStorage и SessionStorage? В каких сценариях следует использовать каждый из них?",
            "LocalStorage хранит данные бессрочно до очистки пользователем или приложением. SessionStorage хранит данные только в пределах текущей вкладки. "
            "LocalStorage подходит для кэша, настроек и черновиков, SessionStorage — для временного состояния интерфейса в текущей сессии.",
        ),
        (
            "2. Как механизм Promise улучшает работу с асинхронными операциями? Какие преимущества предоставляют async/await?",
            "Promise дает единый способ обработки успешного и ошибочного завершения операции. Async/await делает асинхронный код линейным и более читаемым, "
            "упрощает обработку ошибок через try/catch и снижает вероятность «callback hell».",
        ),
        (
            "3. Какие стратегии обработки ошибок вы реализовали при работе с API? Как обеспечить отказоустойчивость приложения?",
            "Реализованы таймауты, retry для временных ошибок, обработка HTTP-статусов, проверка navigator.onLine и локальная очередь отправки. "
            "Отказоустойчивость обеспечивается кэшированием данных, резервным оффлайн-сценарием и последующей синхронизацией.",
        ),
        (
            "4. Какие методы кэширования данных вы применили? Как поддерживать актуальность кэшированных данных?",
            "Использован кэш ленты в LocalStorage с меткой времени и версией схемы хранения. Актуальность поддерживается TTL, перезагрузкой данных с сервера "
            "и принудительным сбросом устаревшего кэша при обновлении версии приложения.",
        ),
        (
            "5. Как реализовать оффлайн-функциональность? Какие данные наиболее критичны для кэширования?",
            "Оффлайн-функциональность строится на локальном кэше и очереди несинхронизированных действий. "
            "Критично кэшировать ленту, черновики, очередь публикаций и пользовательские настройки, чтобы интерфейс оставался работоспособным без сети.",
        ),
        (
            "6. Что такое CORS и как он влияет на взаимодействие с API? Какие методы обхода CORS вы знаете?",
            "CORS — это политика браузера, ограничивающая междоменные запросы. Если сервер не возвращает нужные заголовки, запрос блокируется. "
            "Корректные способы решения: настройка CORS на сервере, прокси на backend, использование того же домена для API и фронтенда.",
        ),
        (
            "7. Каковы лучшие практики безопасности при работе с API ключами на клиенте?",
            "Не размещать секретные ключи напрямую в публичном фронтенде, использовать backend-прокси, ограничивать ключи по доменам и правам, "
            "применять HTTPS, а для пользовательских данных использовать OAuth и токены с ограниченным сроком жизни.",
        ),
        (
            "8. Как реализовать пагинацию при работе с большими объемами данных от API?",
            "Использовать параметры offset/limit или page/pageSize, хранить текущий индекс и подгружать данные пакетами. "
            "В интерфейсе удобно применять бесконечную прокрутку с IntersectionObserver и остановкой при достижении конца данных.",
        ),
        (
            "9. Какие HTTP-статусы наиболее важны при работе с REST API? Как их правильно обрабатывать?",
            "Ключевые статусы: 200/201 (успех), 400 (ошибка клиента), 401/403 (доступ), 404 (не найдено), 429 (лимиты), 500+ (ошибки сервера). "
            "Для каждого класса нужен отдельный сценарий: уведомление пользователю, повтор, логирование или постановка в очередь.",
        ),
        (
            "10. Как оптимизировать производительность при частых запросах к API?",
            "Использовать кэш, debounce/throttle, батчинг, пагинацию, отмену устаревших запросов (AbortController), повтор с ограничением и минимизацию перерисовок UI. "
            "Также полезно загружать только необходимые поля данных и избегать дублирующих запросов.",
        ),
    ]

    for question, answer in qa_items:
        add_paragraph(doc, question, first_indent=False)
        add_paragraph(doc, answer)

    add_heading(doc, "ВЫВОД")
    add_paragraph(
        doc,
        "В ходе лабораторной работы №6 в проекте SocialSphere реализована полноценная асинхронная модель работы с внешним сервисом: "
        "загрузка и публикация данных через Fetch, обработка ошибок, кэширование и синхронизация локального состояния с сервером.",
    )
    add_paragraph(
        doc,
        "Для варианта №5 выполнены ключевые требования: автосохранение черновиков, очередь отправки постов при оффлайн-режиме, "
        "восстановление и синхронизация после появления сети, а также отображение состояний в интерфейсе. "
        "Полученные результаты подтверждают достижение цели работы и готовность проекта к дальнейшей интеграции в React-этапе.",
    )


def export_pdf(doc_path: Path, pdf_path: Path) -> None:
    powershell_script = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$document = $word.Documents.Open('{doc_path}')
$document.SaveAs([ref] '{pdf_path}', [ref] 17)
$document.Close()
$word.Quit()
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", powershell_script], check=True, cwd=BASE_DIR)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    subprocess.run(["python", str(SCRIPT_DIR / "generate_lab6_report_assets.py")], check=True, cwd=BASE_DIR)

    document = configure_document()
    add_title_page(document)
    fill_content(document)
    document.save(FINAL_DOC_PATH)

    try:
        export_pdf(FINAL_DOC_PATH, FINAL_PDF_PATH)
    except Exception as error:
        print(f"PDF export skipped: {error}")

    print(FINAL_DOC_PATH)
    if FINAL_PDF_PATH.exists():
        print(FINAL_PDF_PATH)


if __name__ == "__main__":
    main()
